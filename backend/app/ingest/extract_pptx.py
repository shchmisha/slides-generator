"""PowerPoint: slide titles, text, tables, chart data and speaker notes."""

import io
import logging
from collections.abc import Iterator

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape

from app.ingest.errors import IngestError
from app.ingest.models import Block, Document, FileKind, RawCell, RawTable, RawValue

logger = logging.getLogger(__name__)


def extract_pptx(data: bytes, filename: str) -> Document:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:  # python-pptx raises many types for broken files
        logger.warning("pptx open failed for %s: %r", filename, exc)
        raise IngestError("Не удалось прочитать презентацию: файл повреждён") from exc

    blocks: list[Block] = []
    tables: list[RawTable] = []
    for number, slide in enumerate(presentation.slides, start=1):
        source = f"{filename} › слайд {number}"
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        blocks.append(Block(kind="heading", text=f"Слайд {number}: {title}".rstrip(": "), level=2))
        for shape in _walk(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:  # type: ignore[attr-defined]
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        kind = "list_item" if paragraph.level > 0 else "paragraph"
                        blocks.append(Block(kind=kind, text=text, level=paragraph.level))
            if getattr(shape, "has_table", False):
                tables.append(_table(shape, source, title or None))
            if getattr(shape, "has_chart", False):
                tables.append(_chart(shape, source, title or None))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            text = notes.text.strip() if notes is not None else ""
            if text:
                blocks.append(Block(kind="paragraph", text=f"Заметки к слайду {number}: {text}"))
    return Document(filename=filename, kind=FileKind.PPTX, blocks=blocks, tables=tables, notes=[])


def _walk(shapes: Iterator[BaseShape]) -> Iterator[BaseShape]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP and isinstance(shape, GroupShape):
            yield from _walk(iter(shape.shapes))
        else:
            yield shape


def _table(shape: BaseShape, source: str, title: str | None) -> RawTable:
    rows = [
        [RawCell(value=cell.text.strip() or None) for cell in row.cells]
        for row in shape.table.rows  # type: ignore[attr-defined]
    ]
    return RawTable(source=source, title=title, rows=rows)


def _chart(shape: BaseShape, source: str, title: str | None) -> RawTable:
    """Chart data as a table: categories down, one column per series."""
    chart = shape.chart  # type: ignore[attr-defined]
    plot = chart.plots[0]
    series = list(plot.series)
    header = [RawCell(value="Категория")] + [RawCell(value=s.name) for s in series]
    rows = [header]
    for i, category in enumerate(plot.categories):
        values: list[RawValue] = [s.values[i] if i < len(s.values) else None for s in series]
        rows.append([RawCell(value=str(category))] + [RawCell(value=v) for v in values])
    if chart.has_title and chart.chart_title.has_text_frame:
        title = chart.chart_title.text_frame.text.strip() or title
    return RawTable(source=f"{source} › диаграмма", title=title, rows=rows)
