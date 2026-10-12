"""Excel: every visible sheet is split into blocks separated by empty rows.

A block with one non-empty cell per row is text; anything wider is a table. A single-cell row at
the top of a table block (or a one-line text block right above it) is the table title.
Formulas are never evaluated: the values Excel saved with the file are read.
"""

import datetime as dt
import io
import logging

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from app.ingest.errors import IngestError
from app.ingest.models import Block, Document, FileKind, Note, RawCell, RawTable, RawValue

logger = logging.getLogger(__name__)

MAX_ROWS = 5_000
MAX_COLUMNS = 100

Grid = list[list[RawCell]]


def extract_xlsx(data: bytes, filename: str) -> Document:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for broken files
        logger.warning("xlsx open failed for %s: %r", filename, exc)
        raise IngestError("Не удалось прочитать таблицу Excel: файл повреждён") from exc

    blocks: list[Block] = []
    tables: list[RawTable] = []
    notes: list[Note] = []
    try:
        for sheet in workbook.worksheets:
            source = f"{filename} › {sheet.title}"
            if sheet.sheet_state != "visible":
                notes.append(Note(severity="info", source=source, message="Скрытый лист пропущен"))
                continue
            if not isinstance(sheet, ReadOnlyWorksheet):  # chartsheets have no cells
                continue
            grid, truncated = _read_grid(sheet)
            if truncated:
                notes.append(
                    Note(
                        severity="warning",
                        source=source,
                        message=f"Лист слишком большой: прочитаны первые {MAX_ROWS} строк "
                        f"и {MAX_COLUMNS} столбцов",
                    )
                )
            sheet_blocks, sheet_tables = _split(grid, source)
            blocks += sheet_blocks
            tables += sheet_tables
    finally:
        workbook.close()
    return Document(
        filename=filename, kind=FileKind.XLSX, blocks=blocks, tables=tables, notes=notes
    )


def _coerce(value: object) -> RawValue:
    if value is None or (isinstance(value, (int, float)) and not isinstance(value, bool)):
        return value
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _read_grid(sheet: ReadOnlyWorksheet) -> tuple[Grid, bool]:
    grid: Grid = []
    for r, row in enumerate(
        # The stubs type iter_rows for Worksheet only; ReadOnlyWorksheet reuses the same method.
        sheet.iter_rows(  # type: ignore[misc]
            min_row=1, max_row=MAX_ROWS, min_col=1, max_col=MAX_COLUMNS, values_only=True
        ),
        start=1,
    ):
        grid.append(
            [
                RawCell(value=_coerce(value), address=f"{get_column_letter(c)}{r}")
                for c, value in enumerate(row, start=1)
            ]
        )
    max_row, max_col = sheet.max_row or 0, sheet.max_column or 0
    return grid, max_row > MAX_ROWS or max_col > MAX_COLUMNS


def _filled(row: list[RawCell]) -> list[int]:
    return [i for i, cell in enumerate(row) if cell.value is not None]


def _split(grid: Grid, source: str) -> tuple[list[Block], list[RawTable]]:
    groups: list[Grid] = []
    current: Grid = []
    for row in grid:
        if _filled(row):
            current.append(row)
        elif current:
            groups.append(current)
            current = []
    if current:
        groups.append(current)

    blocks: list[Block] = []
    tables: list[RawTable] = []
    pending_title: str | None = None
    for group in groups:
        first = min(_filled(row)[0] for row in group)
        last = max(_filled(row)[-1] for row in group)
        rows = [row[first : last + 1] for row in group]
        if all(len(_filled(row)) == 1 for row in rows):
            texts = [str(next(c.value for c in row if c.value is not None)) for row in rows]
            blocks += [Block(kind="paragraph", text=text) for text in texts]
            pending_title = texts[-1] if len(texts) == 1 else None
            continue
        title = pending_title
        if len(_filled(rows[0])) == 1 and len(rows) >= 3:
            cell = next(c for c in rows[0] if c.value is not None)
            title, rows = str(cell.value), rows[1:]
        tables.append(RawTable(source=source, title=title, rows=rows))
        pending_title = None
    return blocks, tables
