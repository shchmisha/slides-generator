"""The real pptx-builder skill, called the way the render job calls it. Needs Node and a
Chromium (CHROMIUM_PATH, Playwright's cache, or Debian's package); skipped otherwise."""

from __future__ import annotations

import json
import os
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from app.agent.planner import plan_deck
from app.agent.review import deck_basis, review
from app.agent.settings import DeckSettings
from app.brand import DEFAULT_THEME
from app.config import get_settings
from app.facts.build import build_facts
from app.ingest.extract import extract_file
from app.render import skill
from app.render.skill import SkillError, build_deck

from .fakes import ScriptedModel

FIXTURES = Path(__file__).parent / "fixtures"
SKILL_DIR = Path(__file__).resolve().parents[2] / "skills" / "pptx-builder"


def _chromium() -> str | None:
    if os.environ.get("CHROMIUM_PATH"):
        return os.environ["CHROMIUM_PATH"]
    cache = Path.home() / "Library" / "Caches" / "ms-playwright"
    for candidate in sorted(cache.glob("chromium_headless_shell-*/*/chrome-headless-shell")):
        return str(candidate)
    debian = Path("/usr/lib/chromium/chromium")
    return str(debian) if debian.exists() else None


CHROMIUM = _chromium()
needs_chromium = pytest.mark.skipif(
    shutil.which("node") is None
    or CHROMIUM is None
    or not (SKILL_DIR / "node_modules" / "playwright-core").is_dir(),
    reason="needs Node, Chromium and the skill's node_modules",
)


@needs_chromium
def test_board_report_is_built_by_the_skill(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert CHROMIUM is not None
    monkeypatch.setattr(get_settings(), "pptx_skill_dir", SKILL_DIR)
    monkeypatch.setenv("CHROMIUM_PATH", CHROMIUM)
    board = json.loads((FIXTURES / "board_scenario.json").read_text())
    documents = [
        extract_file((FIXTURES / name).read_bytes(), name)
        for name in ("q3_metrics.xlsx", "manager_comments.docx")
    ]
    store, notes = build_facts(documents)
    plan = plan_deck(documents, store, notes, DeckSettings(), ScriptedModel(json.dumps(board)))
    checked = review(board, deck_basis(plan, DeckSettings()))
    spec = json.dumps(checked.report["preview"], ensure_ascii=False)
    logo = (SKILL_DIR / "examples" / "logo.png").read_bytes()

    result = build_deck(spec, DEFAULT_THEME, logo, tmp_path / "job")

    deck = Presentation(str(result.pptx))
    assert len(deck.slides) == len(result.previews) == checked.report["slides"]
    assert not [i for i in result.issues if i["kind"] in ("overflow", "outside")]
    assert result.timings["layout_ms"] > 0
    texts = " ".join(
        shape.text_frame.text
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "2 877" in texts  # numbers from the table reach the slides


def test_the_skill_gets_no_secrets_and_is_bounded(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(get_settings(), "pptx_skill_dir", SKILL_DIR)
    for name in ("DATABASE_URL", "REDIS_URL", "LLM_API_KEY", "ELEVENLABS_KEY"):
        monkeypatch.setenv(name, "tok-3f9a1c")
    seen: dict[str, object] = {}

    def fake_run(cmd: list[str], cwd: Path, env: dict[str, str], timeout: int) -> tuple[int, str]:
        seen.update(cmd=cmd, env=env, timeout=timeout)
        return 1, "boom"

    monkeypatch.setattr(skill, "_run_group", fake_run)
    with pytest.raises(SkillError, match="Не удалось собрать"):
        build_deck("{}", DEFAULT_THEME, None, tmp_path / "job")
    env = seen["env"]
    assert isinstance(env, dict)
    assert not {"DATABASE_URL", "REDIS_URL", "LLM_API_KEY", "ELEVENLABS_KEY"} & set(env)
    assert "tok-3f9a1c" not in " ".join(env.values())
    # Chromium's profile and sockets: a private folder with a short path, removed afterwards.
    assert len(env["HOME"]) < 40 and env["TMPDIR"] == env["HOME"]
    assert not Path(env["HOME"]).exists()
    assert seen["timeout"] == get_settings().render_timeout_s


def test_font_warnings_follow_the_skills_substitutes(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "pptx_skill_dir", SKILL_DIR)
    installed = frozenset({"liberation sans", "gelasio", "dejavu serif"})
    monkeypatch.setattr(skill, "installed_fonts", lambda: installed)
    assert skill.font_warnings({"fonts": {"heading": "Georgia", "body": "Arial"}}) == []
    warned = skill.font_warnings({"fonts": {"heading": "PT Sans", "body": "Arial"}})
    assert len(warned) == 1 and "PT Sans" in warned[0]
    # Without the Cyrillic stand-in, Georgia is no longer measured safely.
    monkeypatch.setattr(skill, "installed_fonts", lambda: frozenset({"gelasio"}))
    assert skill.font_warnings({"fonts": {"heading": "Georgia", "body": "Georgia"}})
