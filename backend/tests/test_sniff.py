"""Uploads are identified by content; hostile containers are rejected before any parser runs."""

import io
import struct
import time
import zipfile

import docx
import pytest

from app.ingest.errors import IngestError
from app.ingest.models import FileKind
from app.ingest.sniff import MAX_ENTRIES, sniff

DOCM_CT = "application/vnd.ms-word.document.macroEnabled.main+xml"
DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
VBA_REL = "http://schemas.microsoft.com/office/2006/relationships/vbaProject"


def make_docx(text: str = "Привет") -> bytes:
    document = docx.Document()
    document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def rezip(data: bytes, extra: dict[str, bytes] | None = None, edit=None) -> bytes:  # type: ignore[no-untyped-def]
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            payload = src.read(info)
            if edit is not None:
                payload = edit(info.filename, payload)
            dst.writestr(info, payload)
        for name, payload in (extra or {}).items():
            dst.writestr(name, payload)
    return out.getvalue()


def test_docx_is_detected_by_content() -> None:
    result = sniff(make_docx(), "comments.docx")
    assert result.kind is FileKind.DOCX
    assert result.notes == []


def test_exe_renamed_to_docx_is_rejected() -> None:
    with pytest.raises(IngestError, match="не поддерживается"):
        sniff(b"MZ\x90\x00\x03\x00\x00\x00" + b"\x00" * 100, "report.docx")


def test_legacy_or_password_protected_office_is_rejected() -> None:
    with pytest.raises(IngestError, match="Старый формат"):
        sniff(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 100, "old.doc")


def test_plain_zip_that_is_not_office_is_rejected() -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("readme.txt", "hello")
    with pytest.raises(IngestError, match="не документ Office"):
        sniff(out.getvalue(), "fake.xlsx")


def test_wrong_extension_is_read_by_content_with_a_note() -> None:
    result = sniff(make_docx(), "table.xlsx")
    assert result.kind is FileKind.DOCX
    assert "не совпадает" in result.notes[0].message


def test_text_files_need_a_text_extension_and_utf8() -> None:
    assert sniff("Суперсервис — это…".encode(), "superservice.txt").kind is FileKind.TEXT
    with pytest.raises(IngestError):
        sniff(b"\xff\xfe\x00binary", "data.txt")
    with pytest.raises(IngestError):
        sniff(b"plain text", "script.sh")


def test_path_traversal_entry_is_rejected() -> None:
    data = rezip(make_docx(), extra={"../../etc/passwd": b"x"})
    with pytest.raises(IngestError, match="недопустимые пути"):
        sniff(data, "evil.docx")


def test_zip_bomb_is_rejected_by_compression_ratio() -> None:
    data = rezip(make_docx())
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info))
        dst.writestr("word/media/bomb.bin", b"\x00" * (5 * 1024 * 1024), zipfile.ZIP_DEFLATED)
    with pytest.raises(IngestError, match="архивную бомбу"):
        sniff(out.getvalue(), "bomb.docx")


def test_too_many_entries_is_rejected() -> None:
    extra = {f"word/media/{i}.xml": b"<a/>" for i in range(MAX_ENTRIES + 1)}
    with pytest.raises(IngestError, match="слишком много"):
        sniff(rezip(make_docx(), extra=extra), "many.docx")


def test_macros_are_removed_and_the_document_still_opens() -> None:
    def to_docm(name: str, payload: bytes) -> bytes:
        if name == "[Content_Types].xml":
            return payload.replace(DOCX_CT.encode(), DOCM_CT.encode())
        if name == "word/_rels/document.xml.rels":
            rel = f'<Relationship Id="rIdVba" Type="{VBA_REL}" Target="vbaProject.bin"/>'
            return payload.replace(b"</Relationships>", rel.encode() + b"</Relationships>")
        return payload

    docm = rezip(make_docx("С макросом"), extra={"word/vbaProject.bin": b"VBA!"}, edit=to_docm)
    result = sniff(docm, "with_macro.docm")
    assert result.kind is FileKind.DOCX
    assert any("макросы" in note.message for note in result.notes)
    with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
        names = archive.namelist()
        rels = archive.read("word/_rels/document.xml.rels")
    assert "word/vbaProject.bin" not in names
    assert b"vbaProject" not in rels
    assert docx.Document(io.BytesIO(result.data)).paragraphs[0].text == "С макросом"


def test_oversized_central_directory_is_rejected_before_unzipping() -> None:
    # Few entries but very long names: the directory alone is over the limit.
    extra = {f"word/media/{i:04d}" + "x" * 2000 + ".xml": b"<a/>" for i in range(600)}
    with pytest.raises(IngestError, match="слишком много"):
        sniff(rezip(make_docx(), extra=extra), "long_names.docx")


def test_unusual_compression_is_rejected() -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(make_docx())) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info.filename, src.read(info), zipfile.ZIP_BZIP2)
    with pytest.raises(IngestError, match="нестандартным"):
        sniff(out.getvalue(), "bzip.docx")


def test_macro_part_is_found_by_its_content_type_not_only_its_name() -> None:
    vba_ct = "application/vnd.ms-office.vbaProject"

    def declare(name: str, payload: bytes) -> bytes:
        if name == "[Content_Types].xml":
            override = f'<Override PartName="/word/custom.dat" ContentType="{vba_ct}"/>'
            return payload.replace(b"</Types>", override.encode() + b"</Types>")
        if name == "word/_rels/document.xml.rels":
            rel = f'<Relationship Id="rIdVba" Type="{VBA_REL}" Target="custom.dat"/>'
            return payload.replace(b"</Relationships>", rel.encode() + b"</Relationships>")
        return payload

    data = rezip(make_docx(), extra={"word/custom.dat": b"VBA!"}, edit=declare)
    result = sniff(data, "renamed_macro.docx")
    with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
        assert "word/custom.dat" not in archive.namelist()
        assert vba_ct.encode() not in archive.read("[Content_Types].xml")
        assert b"vbaProject" not in archive.read("word/_rels/document.xml.rels")
    assert any("макросы" in note.message for note in result.notes)


def test_damaged_member_is_a_readable_error() -> None:
    docm = rezip(make_docx(), extra={"word/vbaProject.bin": b"VBA!" * 100})
    # Flip a byte inside the compressed document body: its checksum no longer matches.
    with zipfile.ZipFile(io.BytesIO(docm)) as archive:
        info = archive.getinfo("word/document.xml")
    offset = info.header_offset + 30 + len(info.filename) + 10
    broken = docm[:offset] + bytes([docm[offset] ^ 0xFF]) + docm[offset + 1 :]
    with pytest.raises(IngestError, match="повреждён"):
        sniff(broken, "broken.docm")


def with_macro(edit=None) -> bytes:  # type: ignore[no-untyped-def]
    """A docx with a VBA part, so the macro stripper runs."""
    return rezip(make_docx("С макросом"), extra={"word/vbaProject.bin": b"VBA!"}, edit=edit)


@pytest.mark.parametrize(
    ("member", "token"),
    [("[Content_Types].xml", b"<Override "), ("word/_rels/document.xml.rels", b"<Relationship ")],
)
def test_hostile_markup_is_stripped_in_linear_time(member: str, token: bytes) -> None:
    # A comment full of unterminated tags made the old regular expressions run for minutes.
    junk = b"<!--" + token * 40_000 + b"-->"

    def stuff(name: str, payload: bytes) -> bytes:
        return payload.replace(b"?>", b"?>" + junk, 1) if name == member else payload

    started = time.perf_counter()
    result = sniff(with_macro(stuff), "slow.docm")
    assert time.perf_counter() - started < 2
    assert docx.Document(io.BytesIO(result.data)).paragraphs[0].text == "С макросом"


def test_zip64_record_cannot_hide_a_huge_directory() -> None:
    # The plain end record claims one entry; a zip64 record in front of it, which zipfile
    # always prefers, declares ~200 000. The check must read what zipfile reads.
    local = struct.pack("<4s2B4HL2L2H", b"PK\x03\x04", 20, 0, 0, 0, 0, 0, 0, 0, 0, 1, 0) + b"a"
    record = struct.pack("<4s4B4HL2L5H2L", b"PK\x01\x02", 20, 0, 20, 0, *([0] * 14))
    count = 200_000
    directory = record * count
    zip64_at = len(local) + len(directory)
    zip64 = struct.pack(
        "<4sQ2H2L4Q", b"PK\x06\x06", 44, 45, 45, 0, 0, count, count, len(directory), len(local)
    )
    locator = struct.pack("<4sLQL", b"PK\x06\x07", 0, zip64_at, 1)
    end = struct.pack("<4s4H2LH", b"PK\x05\x06", 0, 0, 1, 1, 46, len(local), 0)
    started = time.perf_counter()
    with pytest.raises(IngestError, match="слишком много"):
        sniff(local + directory + zip64 + locator + end, "zip64.docx")
    assert time.perf_counter() - started < 0.5


def test_zip64_and_commented_archives_are_still_accepted(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(zipfile, "ZIP_FILECOUNT_LIMIT", 1)  # forces the zip64 end record
    zip64 = rezip(make_docx())
    monkeypatch.undo()
    with zipfile.ZipFile(io.BytesIO(zip64)) as archive:
        assert archive.infolist()
    assert b"PK\x06\x06" in zip64
    assert sniff(zip64, "zip64.docx").kind is FileKind.DOCX

    out = io.BytesIO(make_docx())
    with zipfile.ZipFile(out, "a") as archive:
        archive.comment = b"saved by an exotic editor"
    assert sniff(out.getvalue(), "comment.docx").kind is FileKind.DOCX


def test_unsupported_zip_version_is_a_readable_error() -> None:
    data = bytearray(make_docx())
    directory = data.find(b"PK\x01\x02")
    data[directory + 6] = 64  # "version needed to extract" 6.4: zipfile refuses it
    with pytest.raises(IngestError, match="повреждён"):
        sniff(bytes(data), "future.docx")


def test_macro_part_declared_by_default_extension_is_removed() -> None:
    vba_ct = "application/vnd.ms-office.vbaProject"

    def declare(name: str, payload: bytes) -> bytes:
        if name == "[Content_Types].xml":
            default = f'<Default Extension="dat" ContentType="{vba_ct}"/>'
            return payload.replace(b"</Types>", default.encode() + b"</Types>")
        return payload

    data = rezip(make_docx(), extra={"word/settings.dat": b"VBA!"}, edit=declare)
    result = sniff(data, "default.docx")
    with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
        assert "word/settings.dat" not in archive.namelist()
        assert vba_ct.encode() not in archive.read("[Content_Types].xml")


@pytest.mark.parametrize(
    "rel",
    [
        "<Relationship Id='rIdVba' Type='{type}' Target='vbaProject.bin'/>",
        '<Relationship Id="rIdVba" Type="{type}" Target="vbaProject.bin"></Relationship>',
        '<Relationship Target="vbaProject.bin" Id="rIdVba" Type="{escaped}"/>',
        '<Relationship Id="rIdVba" Type="http://example.com/innocent" Target="vbaProject.bin"/>',
    ],
)
def test_macro_relationships_are_removed_however_they_are_written(rel: str) -> None:
    escaped = VBA_REL.replace("vba", "&#118;ba")
    markup = rel.format(type=VBA_REL, escaped=escaped).encode()

    def add(name: str, payload: bytes) -> bytes:
        if name == "word/_rels/document.xml.rels":
            return payload.replace(b"</Relationships>", markup + b"</Relationships>")
        return payload

    result = sniff(with_macro(add), "quoted.docm")
    with zipfile.ZipFile(io.BytesIO(result.data)) as archive:
        rels = archive.read("word/_rels/document.xml.rels")
    assert b"rIdVba" not in rels
    assert b"styles.xml" in rels  # the other relationships are kept
    assert docx.Document(io.BytesIO(result.data)).paragraphs[0].text == "С макросом"
