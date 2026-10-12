"""PPTX emitter: re-creates the measured layout as native PowerPoint objects with python-pptx.

Text stays editable text, tables are real tables and charts are real charts with their data
in an embedded workbook ("Edit Data" works). Nothing is rasterised except the logo.
"""

from __future__ import annotations

import base64
import binascii
import io
from pathlib import Path
from typing import TYPE_CHECKING, cast

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu

from pptx_builder.charts import add_chart
from pptx_builder.layout import (
    CellLine,
    ImageElement,
    Layout,
    Para,
    Rgba,
    ShapeElement,
    TableElement,
    TextElement,
)
from pptx_builder.units import emu, pt_from_px, rect

if TYPE_CHECKING:
    from pptx.dml.fill import FillFormat
    from pptx.slide import Slide
    from pptx.table import _Cell
    from pptx.text.text import Font, _Paragraph

# Built-in "No Style, No Grid" table style: every border and fill comes from the layout.
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

ALIGN = {
    "left": PP_ALIGN.LEFT,
    "right": PP_ALIGN.RIGHT,
    "center": PP_ALIGN.CENTER,
    "justify": PP_ALIGN.JUSTIFY,
}
VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
IMAGE_DATA_URL_PREFIXES = ("data:image/png;base64,", "data:image/jpeg;base64,")
BLANK_LAYOUT = 6


class EmitError(ValueError):
    """The layout cannot be turned into a PPTX (e.g. a malformed image)."""


def write_pptx(layout: Layout, out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = emu(layout.slide_size.width_px)
    prs.slide_height = emu(layout.slide_size.height_px)
    prs.core_properties.title = layout.title
    blank = prs.slide_layouts[BLANK_LAYOUT]
    background = RGBColor.from_string(layout.theme.rgb("background"))

    for measured in layout.slides:
        slide = prs.slides.add_slide(blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = background
        for el in measured.elements:
            match el.kind:
                case "text":
                    _add_text(slide, el)
                case "shape":
                    _add_shape(slide, el)
                case "image":
                    _add_image(slide, el)
                case "table":
                    _add_table(slide, el)
                case "chart":
                    add_chart(slide, el, layout.theme)
    prs.save(str(out_path))


# ----------------------------------------------------------------------------------- text


def _style_font(
    font: Font, *, name: str, size_px: float, bold: bool, italic: bool, hex_: str
) -> None:
    font.name = name
    font.size = pt_from_px(size_px)
    font.bold = bold
    font.italic = italic
    font.color.rgb = RGBColor.from_string(hex_)


def _add_text(slide: Slide, el: TextElement) -> None:
    paras = [p for p in el.paras if p.runs]
    if not paras:
        return
    frame = slide.shapes.add_textbox(*rect(el.box)).text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Emu(0)
    frame.vertical_anchor = VALIGN[el.valign]
    for i, para in enumerate(paras):
        _fill_paragraph(frame.paragraphs[0] if i == 0 else frame.add_paragraph(), para)


def _fill_paragraph(paragraph: _Paragraph, para: Para) -> None:
    paragraph.alignment = ALIGN[para.align]
    # Exact line spacing, as in the browser: the measured boxes only fit with the same leading.
    paragraph.line_spacing = pt_from_px(para.line_height_px)
    if para.space_after_px > 0:
        paragraph.space_after = pt_from_px(para.space_after_px)
    for run in para.runs:
        r = paragraph.add_run()
        r.text = run.text
        _style_font(
            r.font,
            name=run.font,
            size_px=run.size_px,
            bold=run.bold,
            italic=run.italic,
            hex_=run.color.hex,
        )


# --------------------------------------------------------------------------------- shapes


def _solid_fill(fill: FillFormat, color: Rgba, owner: etree._Element) -> None:
    """Solid fill with optional transparency (python-pptx has no API for alpha)."""
    if color.alpha <= 0:
        fill.background()
        return
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(color.hex)
    if color.alpha < 1:
        srgb = owner.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")
        if srgb is None:
            raise EmitError("solid fill without a color")
        alpha = etree.SubElement(srgb, qn("a:alpha"))
        alpha.set("val", str(round(color.alpha * 100_000)))


def _add_shape(slide: Slide, el: ShapeElement) -> None:
    if el.ellipse:
        kind = MSO_SHAPE.OVAL
    elif el.radius_px > 0:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE
    else:
        kind = MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, *rect(el.box))
    shorter_side = min(el.box.w, el.box.h)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE and shorter_side > 0:
        # roundRect's adjustment is the corner radius as a fraction of the shorter side.
        shape.adjustments[0] = min(el.radius_px / shorter_side, 0.5)
    _solid_fill(shape.fill, el.fill, shape._element.spPr)
    if el.border and el.border.px > 0:
        shape.line.color.rgb = RGBColor.from_string(el.border.color.hex)
        shape.line.width = pt_from_px(el.border.px)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False  # no theme shadow: the slide design is flat


def _add_image(slide: Slide, el: ImageElement) -> None:
    prefix = next((p for p in IMAGE_DATA_URL_PREFIXES if el.src.startswith(p)), None)
    if prefix is None:
        raise EmitError("images must be PNG or JPEG data: URLs")
    try:
        data = base64.b64decode(el.src[len(prefix) :], validate=True)
    except binascii.Error as exc:
        raise EmitError("image data is not valid base64") from exc
    slide.shapes.add_picture(io.BytesIO(data), *rect(el.box))


# --------------------------------------------------------------------------------- tables


def _border_xml(tag: str, line: CellLine | None) -> etree._Element:
    if line is None or line.px <= 0 or line.color.alpha <= 0:
        xml = f'<{tag} {nsdecls("a")} w="0"><a:noFill/></{tag}>'
    else:
        width = round(line.px * 0.75 * 12_700)
        xml = (
            f'<{tag} {nsdecls("a")} w="{width}" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill><a:srgbClr val="{line.color.hex}"/></a:solidFill>'
            f'<a:prstDash val="solid"/></{tag}>'
        )
    return cast("etree._Element", parse_xml(xml))


def _set_cell_borders(cell: _Cell, top: CellLine, bottom: CellLine) -> None:
    tc_pr: etree._Element = cell._tc.get_or_add_tcPr()
    tags = ("a:lnL", "a:lnR", "a:lnT", "a:lnB")
    for tag in tags:
        for old in tc_pr.findall(qn(tag)):
            tc_pr.remove(old)
    # Schema order: lnL, lnR, lnT, lnB come before the cell fill.
    for i, (tag, line) in enumerate(zip(tags, (None, None, top, bottom), strict=True)):
        tc_pr.insert(i, _border_xml(tag, line))


def _add_table(slide: Slide, el: TableElement) -> None:
    x, y, _, _ = rect(el.box)
    frame = slide.shapes.add_table(
        len(el.cells), len(el.col_w), x, y, emu(sum(el.col_w)), emu(sum(el.row_h))
    )
    table = frame.table
    table.first_row = False
    table.horz_banding = False
    tbl_pr: etree._Element | None = table._tbl.tblPr
    if tbl_pr is None:
        raise EmitError("table without properties")
    style_id = tbl_pr.find(qn("a:tableStyleId"))
    if style_id is None:
        style_id = etree.SubElement(tbl_pr, qn("a:tableStyleId"))
    style_id.text = NO_STYLE_TABLE

    for i, width in enumerate(el.col_w):
        table.columns[i].width = emu(width)
    for i, height in enumerate(el.row_h):
        table.rows[i].height = emu(height)

    for r, cells in enumerate(el.cells):
        for c, measured in enumerate(cells):
            cell = table.cell(r, c)
            top, right, bottom, left = measured.pad_px
            cell.margin_top, cell.margin_right = emu(top), emu(right)
            cell.margin_bottom, cell.margin_left = emu(bottom), emu(left)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame_ = cell.text_frame
            frame_.word_wrap = True
            paragraph = frame_.paragraphs[0]
            paragraph.alignment = ALIGN[measured.align]
            run = paragraph.add_run()
            run.text = measured.text
            _style_font(
                run.font,
                name=measured.font,
                size_px=measured.size_px,
                bold=measured.bold,
                italic=False,
                hex_=measured.color.hex,
            )
            _solid_fill(cell.fill, measured.fill, cell._tc.get_or_add_tcPr())
            _set_cell_borders(cell, measured.border_top, measured.border_bottom)
