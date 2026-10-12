"""Native charts: clustered columns, lines, and column + line combos with a secondary axis.

python-pptx creates one chart type per chart. Combos are built the standard way: create the
chart with every series (so the embedded workbook holds all the data), then move the line
series into their own <c:lineChart> plots, bound to the primary axes or to a new secondary
pair. Styling then goes through the regular python-pptx API on each plot.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, cast

from lxml import etree
from pptx.chart.axis import CategoryAxis, ValueAxis
from pptx.chart.data import CategoryChartData
from pptx.chart.plot import BarPlot
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
    XL_TICK_MARK,
)
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Pt

from pptx_builder.layout import AxisScale, ChartElement, ChartSeries, ChartSpec, Theme
from pptx_builder.units import rect

if TYPE_CHECKING:
    from pptx.chart.axis import _BaseAxis
    from pptx.chart.chart import Chart
    from pptx.chart.datalabel import DataLabels
    from pptx.oxml.chart.chart import CT_PlotArea
    from pptx.shapes.graphfrm import GraphicFrame
    from pptx.slide import Slide
    from pptx.text.text import Font

AXIS_FONT_PT = 9
LABEL_FONT_PT = 8
BAR_GAP_PCT = 40
LINE_WIDTH_PT = 2
MARKER_SIZE = 7
# Bar-only children of <c:ser>; a series moved into a <c:lineChart> must not carry them.
BAR_ONLY_SERIES_TAGS = ("c:invertIfNegative", "c:pictureOptions", "c:shape")
AXIS_TAGS = ("c:catAx", "c:valAx", "c:dateAx", "c:serAx")
SECONDARY_CAT_AX_ID = "50010"
SECONDARY_VAL_AX_ID = "50020"


def add_chart(slide: Slide, element: ChartElement, theme: Theme) -> None:
    spec = element.spec
    bars = [s for s in spec.series if s.kind == "bar"]
    lines = [s for s in spec.series if s.kind == "line"]
    ordered = bars + lines  # workbook / <c:idx> order

    data = CategoryChartData(number_format=native_format(spec.value_format))
    data.categories = spec.categories
    for series in ordered:
        data.add_series(series.name, series.values, number_format=_series_format(spec, series))

    chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED if bars else XL_CHART_TYPE.LINE_MARKERS
    # add_chart returns the graphic frame (its annotation says Chart).
    frame = cast("GraphicFrame", slide.shapes.add_chart(chart_type, *rect(element.box), data))  # type: ignore[arg-type]
    chart = frame.chart

    plot_area = chart._chartSpace.chart.plotArea
    first_plot = plot_area.xCharts[0]
    primary_ids = [str(ax.get("val")) for ax in first_plot.findall(qn("c:axId"))]
    move_primary = [i for i, s in enumerate(ordered) if s.kind == "line" and s.axis == "primary"]
    move_secondary = [i for i, s in enumerate(ordered) if s.axis == "secondary"]
    if bars and move_primary:
        _move_to_line_plot(plot_area, first_plot, move_primary, primary_ids)
    if move_secondary:
        _add_secondary_axes(plot_area)
        _move_to_line_plot(
            plot_area, first_plot, move_secondary, [SECONDARY_CAT_AX_ID, SECONDARY_VAL_AX_ID]
        )

    _style(chart, element, ordered, theme, primary_ids)


def native_format(code: str) -> str:
    """Excel format code as shown on the slide: a "%" suffix gets a no-break space before it,
    the same way the preview formats it (Russian typography: "12,6 %")."""
    return code.replace('"%"', '"\u00a0%"')


def _series_format(spec: ChartSpec, series: ChartSeries) -> str:
    if series.axis == "secondary" and spec.secondary_format:
        return native_format(spec.secondary_format)
    return native_format(spec.value_format)


def _apply_scale(axis: ValueAxis, scale: AxisScale) -> None:
    axis.minimum_scale = scale.min
    axis.maximum_scale = scale.max
    axis.major_unit = scale.step


# ------------------------------------------------------------------------- XML surgery


def _move_to_line_plot(
    plot_area: CT_PlotArea,
    source_plot: etree._Element,
    series_indices: list[int],
    axis_ids: list[str],
) -> None:
    line_plot = parse_xml(
        f'<c:lineChart {nsdecls("c")}><c:grouping val="standard"/><c:varyColors val="0"/>'
        f'<c:marker val="1"/><c:axId val="{axis_ids[0]}"/><c:axId val="{axis_ids[1]}"/>'
        "</c:lineChart>"
    )
    # Plots come before the axes in <c:plotArea>.
    last_plot = plot_area.xCharts[-1]
    last_plot.addnext(line_plot)
    marker = line_plot.find(qn("c:marker"))
    if marker is None:
        raise ValueError("line plot template has no <c:marker>")
    for ser in source_plot.findall(qn("c:ser")):
        idx = ser.find(qn("c:idx"))
        if idx is None or int(idx.get("val", "-1")) not in series_indices:
            continue
        for tag in BAR_ONLY_SERIES_TAGS:
            for child in ser.findall(qn(tag)):
                ser.remove(child)
        marker.addprevious(ser)


def _add_secondary_axes(plot_area: etree._Element) -> None:
    cat_ax = parse_xml(
        f'<c:catAx {nsdecls("c")}><c:axId val="{SECONDARY_CAT_AX_ID}"/>'
        '<c:scaling><c:orientation val="minMax"/></c:scaling><c:delete val="1"/>'
        '<c:axPos val="b"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
        f'<c:tickLblPos val="nextTo"/><c:crossAx val="{SECONDARY_VAL_AX_ID}"/>'
        '<c:crosses val="autoZero"/><c:auto val="1"/><c:lblAlgn val="ctr"/>'
        '<c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>'
    )
    val_ax = parse_xml(
        f'<c:valAx {nsdecls("c")}><c:axId val="{SECONDARY_VAL_AX_ID}"/>'
        '<c:scaling><c:orientation val="minMax"/></c:scaling><c:delete val="0"/>'
        '<c:axPos val="r"/><c:numFmt formatCode="General" sourceLinked="0"/>'
        '<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
        f'<c:crossAx val="{SECONDARY_CAT_AX_ID}"/><c:crosses val="max"/>'
        '<c:crossBetween val="between"/></c:valAx>'
    )
    axes = [child for child in plot_area if child.tag in {qn(t) for t in AXIS_TAGS}]
    axes[-1].addnext(cat_ax)
    cat_ax.addnext(val_ax)


# ------------------------------------------------------------------------------ styling


def _font(font: Font, theme: Theme, size: int, color: str, *, bold: bool = False) -> None:
    font.name = theme.fonts.body
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor.from_string(color)


def _plain_axis(axis: _BaseAxis, theme: Theme, label_color: str) -> None:
    axis.format.line.fill.background()
    axis.major_tick_mark = XL_TICK_MARK.NONE
    axis.minor_tick_mark = XL_TICK_MARK.NONE
    _font(axis.tick_labels.font, theme, AXIS_FONT_PT, label_color)


def _labels(
    labels: DataLabels, theme: Theme, fmt: str, position: XL_LABEL_POSITION, bold: bool
) -> None:
    labels.show_value = True
    labels.number_format = fmt
    labels.number_format_is_linked = False
    labels.position = position
    _font(labels.font, theme, LABEL_FONT_PT, theme.rgb("text"), bold=bold)


def _axis(chart: Chart, tag: str, axis_id: str) -> etree._Element:
    # Chart.value_axis is not usable here: with two value axes python-pptx returns the second.
    found: list[etree._Element] = chart._chartSpace.chart.plotArea.xpath(
        f"{tag}[c:axId/@val='{axis_id}']"
    )
    if not found:
        raise ValueError(f"chart has no {tag} with id {axis_id}")
    return found[0]


def _style(
    chart: Chart,
    element: ChartElement,
    ordered: list[ChartSeries],
    theme: Theme,
    primary_ids: list[str],
) -> None:
    spec = element.spec
    text, muted, grid = theme.rgb("text"), theme.rgb("muted"), theme.rgb("grid")
    chart.has_title = False
    _font(chart.font, theme, AXIS_FONT_PT, text)

    chart.has_legend = True
    legend = chart.legend
    legend.position = XL_LEGEND_POSITION.BOTTOM
    legend.include_in_layout = False
    _font(legend.font, theme, AXIS_FONT_PT, text)

    category_axis = CategoryAxis(_axis(chart, "c:catAx", primary_ids[0]))
    _plain_axis(category_axis, theme, text)
    category_axis.has_major_gridlines = False

    value_axis = ValueAxis(_axis(chart, "c:valAx", primary_ids[1]))
    _plain_axis(value_axis, theme, muted)
    value_axis.has_major_gridlines = True
    gridlines = value_axis.major_gridlines.format.line
    gridlines.color.rgb = RGBColor.from_string(grid)
    gridlines.width = Pt(0.75)
    value_axis.tick_labels.number_format = native_format(spec.value_format)
    value_axis.tick_labels.number_format_is_linked = False
    _apply_scale(value_axis, element.scales.primary)

    if any(s.axis == "secondary" for s in ordered):
        axis = ValueAxis(_axis(chart, "c:valAx", SECONDARY_VAL_AX_ID))
        _plain_axis(axis, theme, muted)
        axis.has_major_gridlines = False
        axis.tick_labels.number_format = native_format(spec.secondary_format or "General")
        axis.tick_labels.number_format_is_linked = False
        if element.scales.secondary is not None:
            _apply_scale(axis, element.scales.secondary)

    for plot in chart.plots:
        if isinstance(plot, BarPlot):
            plot.gap_width = BAR_GAP_PCT
            plot.overlap = 0
        plot.vary_by_categories = False
        for series in plot.series:
            measured = ordered[series.index]
            color = RGBColor.from_string(theme.rgb(measured.color))
            fmt = _series_format(spec, measured)
            if measured.kind == "bar":
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color
                series.invert_if_negative = False
                if measured.show_values:
                    _labels(series.data_labels, theme, fmt, XL_LABEL_POSITION.OUTSIDE_END, False)
            else:
                line = series.format.line
                line.color.rgb = color
                line.width = Pt(LINE_WIDTH_PT)
                series.smooth = False
                series.marker.style = XL_MARKER_STYLE.CIRCLE
                series.marker.size = MARKER_SIZE
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = color
                series.marker.format.line.color.rgb = RGBColor.from_string("FFFFFF")
                if measured.show_values:
                    _labels(series.data_labels, theme, fmt, XL_LABEL_POSITION.ABOVE, True)

    _transparent_chart_area(chart)


def _transparent_chart_area(chart: Chart) -> None:
    """No fill and no border around the chart, so it sits on the slide like the preview."""
    space = chart._chartSpace
    for old in space.findall(qn("c:spPr")):
        space.remove(old)
    sp_pr = parse_xml(f"<c:spPr {nsdecls('c', 'a')}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>")
    # Schema order: ... c:chart, c:spPr, c:txPr, c:externalData ...
    space.chart.addnext(sp_pr)
