"""End to end: spec → Node/Chromium layout → python-pptx. Needs node and a Chromium build."""

from __future__ import annotations

import json
import os
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from pptx_builder.build import BuildError, build_deck, main

from .conftest import SKILL_DIR

EXAMPLES = SKILL_DIR / "examples"
THEME = SKILL_DIR / "themes" / "brand.json"

needs_browser = pytest.mark.skipif(
    shutil.which("node") is None or not (SKILL_DIR / "node_modules").exists(),
    reason="node and the skill's node_modules are required",
)


@needs_browser
def test_example_deck_builds_without_layout_issues(tmp_path: Path) -> None:
    report = build_deck(EXAMPLES / "q3_deck.json", THEME, tmp_path, EXAMPLES / "logo.png")
    assert report.issues == []
    assert (tmp_path / report.pptx).exists()
    assert all((tmp_path / name).exists() for name in report.previews)

    prs = Presentation(str(tmp_path / report.pptx))
    spec = json.loads((EXAMPLES / "q3_deck.json").read_text())
    assert len(prs.slides) == len(spec["slides"])
    charts = [s.chart for slide in prs.slides for s in slide.shapes if s.has_chart]
    tables = [s.table for slide in prs.slides for s in slide.shapes if s.has_table]
    assert len(charts) == 3  # chart_table ×2, chart
    assert len(tables) == 3  # chart_table ×2, table
    # Numbers on the slide come straight from the spec.
    assert list(charts[0].plots[0].series[1].values) == [379, 674, 1065, 449, 310]


@needs_browser
def test_layout_overflow_is_reported_not_fatal(tmp_path: Path) -> None:
    spec = json.loads((EXAMPLES / "q3_deck.json").read_text())
    spec["slides"] = [spec["slides"][1]]
    spec["slides"][0]["callout"] = "Очень длинный вывод. " * 60
    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps(spec, ensure_ascii=False))
    report = build_deck(spec_path, THEME, tmp_path / "out")
    kinds = {issue["kind"] for issue in report.issues}
    assert kinds & {"overflow", "shrunk"}
    assert (tmp_path / "out" / report.pptx).exists()


@needs_browser
def test_invalid_theme_fails_with_a_readable_error(tmp_path: Path) -> None:
    bad_theme = tmp_path / "theme.json"
    bad_theme.write_text(json.dumps({"colors": {"primary": "red"}, "fonts": {}}))
    with pytest.raises(BuildError, match=r"theme\.colors"):
        build_deck(EXAMPLES / "q3_deck.json", bad_theme, tmp_path / "out")


def test_layout_step_gets_no_secrets(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    seen: dict[str, str] = {}

    def fake_run(cmd: list[str], **kwargs: object) -> object:
        env = kwargs["env"]
        assert isinstance(env, dict)
        seen.update(env)
        raise BuildError("stop here")

    monkeypatch.setenv("LLM_API_KEY", "sk-secret")
    monkeypatch.setenv("DATABASE_URL", "postgresql://u:p@db/x")
    monkeypatch.setattr("pptx_builder.build.subprocess.run", fake_run)
    with pytest.raises(BuildError):
        build_deck(EXAMPLES / "q3_deck.json", THEME, tmp_path)
    assert "LLM_API_KEY" not in seen
    assert "DATABASE_URL" not in seen
    assert seen.get("PATH") == os.environ["PATH"]


def test_cli_requires_spec_and_theme_or_layout(tmp_path: Path) -> None:
    with pytest.raises(SystemExit):
        main(["--out", str(tmp_path)])
