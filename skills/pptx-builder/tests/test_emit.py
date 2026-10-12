"""The emitter turns measured elements into native, editable PowerPoint objects."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as PresentationType

from pptx_builder.emit import NO_STYLE_TABLE, EmitError, write_pptx
from pptx_builder.layout import Layout

from .conftest import cell, chart_el, make_layout, rgba, text_el

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def roundtrip(layout: Layout, tmp_path: Path) -> PresentationType:
    out = tmp_path / "deck.pptx"
    write_pptx(layout, out)
    return Presentation(str(out))


def test_slide_size_is_16_9_widescreen(tmp_path: Path) -> None:
    prs = roundtrip(make_layout([text_el("Заголовок")]), tmp_path)
    assert (prs.slide_width, prs.slide_height) == (12_192_000, 6_858_000)


def test_text_keeps_runs_fonts_and_position(tmp_path: Path) -> None:
    prs = roundtrip(make_layout([text_el("Выручка 2 877 млн ₽", bold=True)]), tmp_path)
    shape = prs.slides[0].shapes[0]
    assert shape.has_text_frame
    run = shape.text_frame.paragraphs[0].runs[0]
    assert run.text == "Выручка 2 877 млн ₽"
    assert run.font.name == "Georgia"
    assert run.font.size.pt == 24  # 32 px
    assert run.font.bold is True
    assert str(run.font.color.rgb) == "1F3A5F"
    assert (shape.left, shape.top) == (64 * 9525, 40 * 9525)


def test_empty_text_is_skipped(tmp_path: Path) -> None:
    element = text_el("x")
    element["paras"][0]["runs"] = []
    prs = roundtrip(make_layout([element]), tmp_path)
    assert len(prs.slides[0].shapes) == 0


def test_shapes_map_to_rect_rounded_rect_and_oval_with_alpha(tmp_path: Path) -> None:
    box = {"x": 10, "y": 10, "w": 100, "h": 50}
    shapes = [
        {"kind": "shape", "box": box, "fill": rgba("F3F5F8"), "radiusPx": 0, "ellipse": False,
         "border": None},
        {"kind": "shape", "box": box, "fill": rgba("1F3A5F", 0.5), "radiusPx": 10,
         "ellipse": False, "border": None},
        {"kind": "shape", "box": {"x": 0, "y": 0, "w": 30, "h": 30}, "fill": rgba("1F3A5F"),
         "radiusPx": 15, "ellipse": True, "border": None},
    ]  # fmt: skip
    prs = roundtrip(make_layout(shapes), tmp_path)
    rect, rounded, oval = prs.slides[0].shapes
    assert rect.auto_shape_type == MSO_SHAPE.RECTANGLE
    assert rounded.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
    assert rounded.adjustments[0] == pytest.approx(0.2)  # 10 px of the 50 px side
    assert oval.auto_shape_type == MSO_SHAPE.OVAL
    alpha = rounded._element.spPr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}/{qn('a:alpha')}")
    assert alpha is not None
    assert alpha.get("val") == "50000"
    assert rect._element.spPr.find(f"{qn('a:ln')}/{qn('a:noFill')}") is not None


def test_logo_becomes_a_picture(tmp_path: Path, logo_data_url: str) -> None:
    image = {"kind": "image", "box": {"x": 1096, "y": 40, "w": 120, "h": 40}, "src": logo_data_url}
    prs = roundtrip(make_layout([image]), tmp_path)
    picture = prs.slides[0].shapes[0]
    assert picture.shape_type == MSO_SHAPE_TYPE.PICTURE
    assert picture.width == 120 * 9525


def test_non_image_data_url_is_rejected(tmp_path: Path) -> None:
    image = {"kind": "image", "box": {"x": 0, "y": 0, "w": 10, "h": 10},
             "src": "data:text/html;base64,PHNjcmlwdD4="}  # fmt: skip
    with pytest.raises(EmitError):
        write_pptx(make_layout([image]), tmp_path / "deck.pptx")


def test_table_is_native_with_measured_sizes_and_borders(tmp_path: Path) -> None:
    table = {
        "kind": "table",
        "box": {"x": 750, "y": 140, "w": 400, "h": 72},
        "colW": [200, 200],
        "rowH": [36, 36],
        "cells": [
            [cell("Филиал", bottom_px=2), cell("Факт", bottom_px=2)],
            [cell("Итого", bottom_px=0, fill="F3F5F8", alpha=1), cell("2 877", bottom_px=0)],
        ],
    }
    prs = roundtrip(make_layout([table]), tmp_path)
    frame = prs.slides[0].shapes[0]
    assert frame.has_table
    tbl = frame.table
    assert [c.width for c in tbl.columns] == [200 * 9525, 200 * 9525]
    assert tbl.cell(1, 1).text == "2 877"
    assert tbl._tbl.tblPr.find(qn("a:tableStyleId")).text == NO_STYLE_TABLE

    header_pr = tbl.cell(0, 0)._tc.tcPr
    tags = [child.tag for child in header_pr]
    # Borders must precede the fill element, or PowerPoint reports the file as damaged.
    assert tags[:4] == [qn("a:lnL"), qn("a:lnR"), qn("a:lnT"), qn("a:lnB")]
    assert header_pr.find(qn("a:lnB")).get("w") == str(round(2 * 0.75 * 12_700))
    assert header_pr.find(qn("a:lnL")).find(qn("a:noFill")) is not None
    total_pr = tbl.cell(1, 0)._tc.tcPr
    assert total_pr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}").get("val") == "F3F5F8"


def test_bar_chart_is_native_with_data_and_shared_scale(tmp_path: Path) -> None:
    chart_element = chart_el(
        [
            {"name": "План", "kind": "bar", "color": "grid", "values": [915, 945, 975]},
            {"name": "Факт", "kind": "bar", "color": "primary", "values": [937, 942, 998],
             "showValues": True},
        ]
    )  # fmt: skip
    prs = roundtrip(make_layout([chart_element]), tmp_path)
    chart = prs.slides[0].shapes[0].chart
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    assert [s.name for s in chart.plots[0].series] == ["План", "Факт"]
    assert list(chart.plots[0].series[1].values) == [937, 942, 998]
    assert chart.plots[0].gap_width == 40
    assert (chart.value_axis.minimum_scale, chart.value_axis.maximum_scale) == (0, 1000)
    assert chart.value_axis.major_unit == 200
    assert chart.part.chart_workbook.xlsx_part is not None  # "Edit Data" works


def test_combo_chart_puts_lines_on_a_secondary_axis(tmp_path: Path) -> None:
    chart_element = chart_el(
        [
            {"name": "Выручка", "kind": "bar", "color": "primary", "values": [937, 942, 998]},
            {"name": "Маржа, %", "kind": "line", "axis": "secondary", "color": "accent",
             "values": [13.2, 11.1, 13.3], "showValues": True},
        ],
        secondaryFormat='0.0"%"',
    )  # fmt: skip
    prs = roundtrip(make_layout([chart_element]), tmp_path)
    chart = prs.slides[0].shapes[0].chart
    plot_area = chart._chartSpace.chart.plotArea
    bar_plot, line_plot = plot_area.xCharts
    assert bar_plot.tag == f"{C}barChart"
    assert line_plot.tag == f"{C}lineChart"

    line_ser = line_plot.find(f"{C}ser")
    assert line_ser.find(f"{C}idx").get("val") == "1"
    assert line_ser.find(f"{C}invertIfNegative") is None  # bar-only element removed

    line_axis_ids = [a.get("val") for a in line_plot.findall(f"{C}axId")]
    secondary = plot_area.xpath(f"c:valAx[c:axId/@val='{line_axis_ids[1]}']")[0]
    assert secondary.find(f"{C}axPos").get("val") == "r"
    assert secondary.find(f"{C}crosses").get("val") == "max"
    assert secondary.find(f"{C}numFmt").get("formatCode") == '0.0" %"'
    assert secondary.find(f"{C}scaling/{C}max").get("val") == "15.0"
    hidden_cat = plot_area.xpath(f"c:catAx[c:axId/@val='{line_axis_ids[0]}']")[0]
    assert hidden_cat.find(f"{C}delete").get("val") == "1"
    # Primary axis keeps the primary scale, not the secondary one.
    primary = plot_area.xpath(f"c:valAx[c:axId/@val!='{line_axis_ids[1]}']")[0]
    assert primary.find(f"{C}scaling/{C}max").get("val") == "1000.0"


def test_chart_area_is_transparent(tmp_path: Path) -> None:
    series = [{"name": "Факт", "kind": "bar", "color": "primary", "values": [1, 2, 3]}]
    prs = roundtrip(make_layout([chart_el(series)]), tmp_path)
    space = prs.slides[0].shapes[0].chart._chartSpace
    sp_pr = space.find(f"{C}spPr")
    assert sp_pr is not None
    assert sp_pr.find(qn("a:noFill")) is not None
    children = [child.tag for child in space]
    assert children.index(f"{C}spPr") == children.index(f"{C}chart") + 1
